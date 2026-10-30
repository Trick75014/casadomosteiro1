"""
rappel.py — Rappels avant l'arrivée des voyageurs (Casa do Mosteiro).
Lit l'Excel du dépôt et ouvre une « issue » GitHub : GitHub t'envoie alors
lui-même un e-mail (et une notification sur l'appli mobile GitHub).
Aucun service externe, aucun mot de passe, aucun secret à créer.

Jours de rappel : variable JOURS (par défaut "5,2" = J-5 et J-2).
Test sans rien créer :  python rappel.py --dry-run --today 2026-11-03
"""
import argparse, glob, os, subprocess, unicodedata
from datetime import datetime, date, timedelta
from zoneinfo import ZoneInfo
from openpyxl import load_workbook

JOURS = [int(x) for x in os.environ.get("JOURS", "5,2").split(",") if x.strip()]
UNCONFIRMED = {"", "a definir", "tarde", "manha", "noite", "none"}


def norm(s):
    s = str(s or "").strip().lower()
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def load_reservations():
    xlsx = sorted(glob.glob("*.xlsx"), key=os.path.getmtime, reverse=True)[0]
    wb = load_workbook(xlsx, data_only=True)
    ws = next((wb[n] for n in wb.sheetnames if "serva" in n.lower()), wb.active)
    hr = next(r for r in range(1, 8) if "voyageur" in " ".join(norm(c.value) for c in ws[r]))
    cols = {norm(c.value): i for i, c in enumerate(ws[hr]) if c.value}

    def col(*keys):
        for k in keys:
            for h, i in cols.items():
                if norm(k) in h:
                    return i

    ix = dict(plat=col("plateforme"), name=col("voyageur"), arr=col("arrivee"), dep=col("depart"),
              nights=col("nuits"), pers=col("personnes"), hour=col("hora"), country=col("origem"),
              fatura=col("fatura"), decl=col("declara"), modelo=col("modelo"), siba=col("siba"))
    out = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        g = lambda k: row[ix[k]] if ix[k] is not None and ix[k] < len(row) else None
        if not g("name") or not isinstance(g("arr"), (datetime, date)):
            continue
        d = lambda v: v.date() if isinstance(v, datetime) else v
        out.append(dict(
            plat=g("plat"), name=str(g("name")).strip(), arr=d(g("arr")), dep=d(g("dep")),
            nights=g("nights"), pers=g("pers"), hour=str(g("hour") or "").strip(),
            country=str(g("country") or "").strip(),
            siba=bool(g("siba") and "siba" in norm(g("siba"))),
            todo=[l for l, k in (("Fatura/Recibo", "fatura"), ("Declaração periódica", "decl"),
                                 ("Modelo 30", "modelo")) if norm(g(k)) != "done"]))
    return out


def build(bookings, days):
    arr = bookings[0]["arr"]
    title = f"🏠 Arrivée {arr:%d/%m/%Y} (dans {days} j) — " + ", ".join(b["name"] for b in bookings)
    parts = []
    for b in bookings:
        ok = norm(b["hour"]) not in UNCONFIRMED
        hour = b["hour"] if ok else f"{b['hour'] or 'non renseignée'} ⚠️ à confirmer avec le voyageur"
        lines = [f"### {b['name']}",
                 f"- **Séjour** : {b['arr']:%d/%m/%Y} → {b['dep']:%d/%m/%Y} ({b['nights']} nuits)",
                 f"- **Personnes** : {b['pers']}",
                 f"- **Heure d'arrivée** : {hour}",
                 f"- **Pays** : {b['country'] or '—'}",
                 f"- **Plateforme** : {b['plat']}"]
        if b["siba"]:
            lines.append("- ⚠️ **SIBA** : voyageur étranger → à déclarer")
        if b["todo"]:
            lines += ["", "**À faire :**"] + [f"- [ ] {t}" for t in b["todo"]]
        lines += ["", "_Téléphone : voir le calendrier ou l'Excel._"]
        parts.append("\n".join(lines))
    return title, "\n\n---\n\n".join(parts)


def create_issue(title, body):
    q = subprocess.run(["gh", "issue", "list", "--state", "all", "--search", f'"{title}" in:title',
                        "--json", "title", "--limit", "20"], capture_output=True, text=True)
    if q.returncode == 0 and f'"title":{__import__("json").dumps(title, ensure_ascii=False)}' in q.stdout.replace(" ", ""):
        print("Déjà créé, on passe :", title)
        return
    cmd = ["gh", "issue", "create", "--title", title, "--body-file", "-"]
    if os.environ.get("OWNER"):
        cmd += ["--assignee", os.environ["OWNER"]]
    r = subprocess.run(cmd, input=body, text=True, capture_output=True)
    print(r.stdout or r.stderr)
    r.check_returncode()


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--today")
    a = ap.parse_args()
    today = date.fromisoformat(a.today) if a.today else datetime.now(ZoneInfo("Europe/Lisbon")).date()
    res = load_reservations()
    found = False
    for d in JOURS:
        todays = [b for b in res if b["arr"] == today + timedelta(days=d)]
        if not todays:
            continue
        found = True
        title, body = build(todays, d)
        print(title)
        print(body if a.dry_run else "")
        if not a.dry_run:
            create_issue(title, body)
    if not found:
        print(f"Aucune arrivée à J-{' ni J-'.join(map(str, JOURS))} ({today})")
